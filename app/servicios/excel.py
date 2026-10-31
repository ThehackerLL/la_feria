"""Exportar los reportes a Excel (carpeta «reportes»).

El Excel es una FOTO del periodo: muestra lo mismo que la pantalla Reportes.
Los costos y la ganancia solo se incluyen si `con_costos=True` (la pantalla lo
pide cuando la ganancia está visible).

Hojas: Resumen, Día por día (o Mes por mes), Ventas, Más vendidos,
Productos parados y Cierres de caja.
"""
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app import config
from app.servicios import caja, reportes, respaldos

FUENTE = "Arial"
MORADO, GRIS = "6A0080", "E5E7EB"
DINERO = '"S/" #,##0.00'
_MESES = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio",
          "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre"]
_ESTADO_CAJA = {"abierta": "Abierta", "cerrada": "Cerrada", "sin_contar": "Sin contar"}
_BORDE = Border(bottom=Side(style="thin", color="D9DDE3"))


def _dmy(texto):
    d = str(texto)[:10].split("-")
    return f"{d[2]}/{d[1]}/{d[0]}"


def _titulo(ws, texto, sub=""):
    ws["A1"] = texto
    ws["A1"].font = Font(name=FUENTE, size=14, bold=True, color=MORADO)
    if sub:
        ws["A2"] = sub
        ws["A2"].font = Font(name=FUENTE, size=10, color="5B6675")


def _tabla(ws, fila, encabezados, filas, formatos=(), anchos=()):
    """Escribe encabezado + filas desde `fila`. formatos: número por columna
    (None = texto). Devuelve la fila siguiente a la última escrita."""
    for j, texto in enumerate(encabezados, 1):
        c = ws.cell(row=fila, column=j, value=texto)
        c.font = Font(name=FUENTE, size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=MORADO)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, valores in enumerate(filas, fila + 1):
        for j, v in enumerate(valores, 1):
            c = ws.cell(row=i, column=j, value=v)
            c.font = Font(name=FUENTE, size=10)
            c.border = _BORDE
            fmt = formatos[j - 1] if j - 1 < len(formatos) else None
            if fmt:
                c.number_format = fmt
            if isinstance(v, (int, float)):
                c.alignment = Alignment(horizontal="right")
            else:
                c.alignment = Alignment(horizontal="left", indent=1)
    for j, ancho in enumerate(anchos, 1):
        ws.column_dimensions[get_column_letter(j)].width = ancho
    return fila + 1 + len(filas)


def _negrita_fila(ws, fila, columnas):
    for j in range(1, columnas + 1):
        c = ws.cell(row=fila, column=j)
        c.font = Font(name=FUENTE, size=10, bold=True)
        c.fill = PatternFill("solid", fgColor=GRIS)


def _pagina(ws):
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def _sin_datos(ws, fila, texto):
    ws.cell(row=fila, column=1, value=texto).font = Font(
        name=FUENTE, size=10, italic=True, color="5B6675")


def _cambio(valor):
    return None if valor is None else valor / 100


# ------------------------------------------------------------------ hojas
def _hoja_resumen(wb, r, con_costos):
    ws = wb.active
    ws.title = "Resumen"
    _titulo(ws, f"Reporte: {r['etiqueta'].lower()}",
            f"Del {_dmy(r['desde'])} al {_dmy(r['hasta'])}   ·   generado "
            f"{datetime.now():%d/%m/%Y %H:%M}")
    antes = r["frente_a"].capitalize()
    filas = [("Total vendido", r["vendido"], r["anterior"]["vendido"], _cambio(r["cambio"]["vendido"])),
             ("Número de ventas", r["ventas"], r["anterior"]["ventas"], _cambio(r["cambio"]["ventas"]))]
    if con_costos:
        filas.append(("Ganancia bruta", r["ganancia"], r["anterior"]["ganancia"],
                      _cambio(r["cambio"]["ganancia"])))
    fin = _tabla(ws, 4, ["Comparación", "Ahora", antes, "Cambio"], filas,
                 [None, DINERO, DINERO, "+0.0%;-0.0%;0.0%"], [54, 18, 18, 14])
    ws.cell(row=5 + 1, column=2).number_format = "0"      # número de ventas, sin S/
    ws.cell(row=5 + 1, column=3).number_format = "0"
    dinero = [("Efectivo recibido", r["efectivo_cobrado"]),
              ("Cobros digitales (Yape, Plin, transferencia, tarjeta)", r["digital_cobrado"]),
              ("Total cobrado", r["cobrado_total"]),
              ("   de eso: abonos y cuotas de clientes", r["abonos_cobrados"]),
              ("Pagos efectuados (compras y cuotas a proveedores)", r["pagado_total"]),
              ("   de eso: cuotas a proveedores", r["cuotas_pagadas"])]
    fin = _tabla(ws, fin + 1, ["Dinero cobrado y pagado", "Monto"], dinero, [None, DINERO])
    ws.cell(row=fin + 1, column=1,
            value=("Este reporte incluye costos y ganancia." if con_costos else
                   "Este reporte no incluye costos ni ganancia.")
            ).font = Font(name=FUENTE, size=9, italic=True, color="5B6675")
    _pagina(ws)


def _hoja_dias(wb, periodo, con_costos, hoy):
    if periodo == "año":
        ws = wb.create_sheet("Mes por mes")
        _titulo(ws, "Mes por mes")
        filas = [(_MESES[m["mes"] - 1], m["ventas"], m["vendido"], m["pagado"]
                  ) + ((m["ganancia"],) if con_costos else ())
                 for m in reportes.meses_del_año(_fecha(hoy).year)[:_fecha(hoy).month]]
        cab = ["Mes", "Ventas", "Vendido", "Compras pagadas"] + (["Ganancia"] if con_costos else [])
    else:
        ws = wb.create_sheet("Día por día")
        _titulo(ws, "Día por día")
        filas = [(_dmy(d["fecha"]), d["ventas"], d["vendido"]) + ((d["ganancia"],) if con_costos else ())
                 for d in reportes.dias_del_periodo(periodo, hoy)]
        cab = ["Fecha", "Ventas", "Vendido"] + (["Ganancia"] if con_costos else [])
    fmts = [None, "0"] + [DINERO] * (len(cab) - 2)
    fin = _tabla(ws, 3, cab, filas, fmts, [18, 10, 16, 18, 16])
    total = ["Total", sum(f[1] for f in filas)] + [round(sum(f[j] for f in filas), 2)
                                                  for j in range(2, len(cab))]
    for j, v in enumerate(total, 1):
        ws.cell(row=fin, column=j, value=v).number_format = fmts[j - 1] or "General"
    _negrita_fila(ws, fin, len(cab))
    ws.freeze_panes = "A4"
    _pagina(ws)


def _hoja_ventas(wb, periodo, con_costos, hoy):
    ws = wb.create_sheet("Ventas")
    _titulo(ws, "Ventas del periodo", "Pagado = lo cobrado hasta hoy; Por cobrar = lo que falta (créditos).")
    ventas = reportes.ventas_del_periodo(periodo, hoy)
    cab = ["Venta", "Fecha", "Hora", "Cliente", "Total", "Pagado", "Por cobrar", "Forma de pago"]
    if con_costos:
        cab += ["Costo", "Ganancia"]
    filas = []
    for v in ventas:
        fila = [v["numero"], _dmy(v["fecha"]), str(v["fecha"])[11:16], v["cliente"],
                v["total"], v["pagado"], v["por_cobrar"], v["formas"]]
        if con_costos:
            fila += [v["costo"], v["ganancia"]]
        filas.append(fila)
    fmts = [None, None, None, None, DINERO, DINERO, DINERO, None, DINERO, DINERO]
    fin = _tabla(ws, 4, cab, filas, fmts, [12, 12, 8, 30, 14, 14, 14, 24, 14, 14])
    if filas:
        ws.cell(row=fin, column=1, value="Total")
        for j in (5, 6, 7) + ((9, 10) if con_costos else ()):
            c = ws.cell(row=fin, column=j, value=round(sum(f[j - 1] for f in filas), 2))
            c.number_format = DINERO
        _negrita_fila(ws, fin, len(cab))
    ws.freeze_panes = "A5"
    _pagina(ws)


def _hoja_productos(wb, periodo, con_costos, hoy):
    ws = wb.create_sheet("Más vendidos")
    _titulo(ws, "Los 10 productos más vendidos")
    mas = reportes.mas_vendidos(periodo, hoy)
    _tabla(ws, 3, ["Producto", "Unidades", "Vendido", "Stock actual"],
           [(f["nombre"], f["unidades"], f["vendido"], f["stock"]) for f in mas],
           [None, "0", DINERO, "0"], [34, 12, 16, 14])
    _pagina(ws)
    ws = wb.create_sheet("Productos parados")
    _titulo(ws, "Productos parados",
            "Tienen stock y casi no se venden. Con «Hoy» se usa la semana.")
    parados = reportes.productos_parados(periodo, hoy)
    cab = ["Producto", "Stock", "Vendió", "Última venta", "Situación"] + (["Dinero parado"] if con_costos else [])
    filas = []
    for f in parados:
        ult = ("Nunca" if f["dias_sin_vender"] is None else _dmy(f["ultima_venta"]))
        fila = [f["nombre"], f["stock"], f["vendido"], ult,
                "No vendió nada" if f["estado"] == "sin_ventas" else "Vendió muy poco"]
        if con_costos:
            fila.append(f["valor_parado"])
        filas.append(fila)
    fin = _tabla(ws, 4, cab, filas, [None, "0", "0", None, None, DINERO], [34, 10, 10, 16, 20, 16])
    if not filas:
        _sin_datos(ws, fin, "Ningún producto parado en este periodo.")
    _pagina(ws)


def _hoja_cierres(wb, periodo, hoy):
    desde, hasta, _, _ = reportes.rango(periodo, hoy)
    ws = wb.create_sheet("Cierres de caja")
    _titulo(ws, "Aperturas y cierres de caja")
    filas = [(_dmy(h["fecha"]), _ESTADO_CAJA.get(h["estado"], h["estado"]), h["apertura"],
              h["esperado"], h["contado"], h["diferencia"], h["nota"])
             for h in caja.historial(400)
             if desde.isoformat() <= h["fecha"] <= hasta.isoformat()]
    fin = _tabla(ws, 3, ["Fecha", "Estado", "Apertura", "Esperado", "Contado", "Diferencia", "Nota"],
                 filas, [None, None, DINERO, DINERO, DINERO, DINERO, None],
                 [12, 12, 14, 14, 14, 14, 36])
    if not filas:
        _sin_datos(ws, fin, "No hubo aperturas ni cierres de caja en este periodo.")
    _pagina(ws)


def _fecha(hoy):
    return hoy if isinstance(hoy, date) else date.today()


# ------------------------------------------------------------------ principal
def carpeta_de_salida(carpeta=None):
    """Dónde se guardará el Excel.

    Si hay una carpeta elegida en Configuración (p. ej. Google Drive) y está
    disponible, ahí; si no está disponible, la carpeta «reportes» de la PC.
    """
    if carpeta:
        return carpeta
    elegida = respaldos.carpeta_excel()
    if elegida is not None and elegida.exists():
        return elegida
    return config.REPORTES_DIR


def exportar_reporte(periodo, con_costos=False, hoy=None, carpeta=None):
    """Crea el Excel del periodo y devuelve su ruta (Path).

    Se guarda en la carpeta elegida en Configuración (Google Drive) si está
    disponible; si no, en «reportes». No se hacen dos copias.

    Si el archivo del mismo día ya existe y está abierto en Excel, guarda uno
    nuevo con un número al final (Reporte_mes_2026-10-22_2.xlsx).
    Con costos el nombre termina en _con_costos..
    """
    hoy = _fecha(hoy)
    carpeta = carpeta_de_salida(carpeta)
    carpeta.mkdir(parents=True, exist_ok=True)
    r = reportes.resumen_periodo(periodo, hoy)
    wb = Workbook()
    _hoja_resumen(wb, r, con_costos)
    _hoja_dias(wb, periodo, con_costos, hoy)
    _hoja_ventas(wb, periodo, con_costos, hoy)
    _hoja_productos(wb, periodo, con_costos, hoy)
    _hoja_cierres(wb, periodo, hoy)
    nombre = {"año": "anio"}.get(periodo, periodo)   # nombre de archivo sin ñ
    base = f"Reporte_{nombre}_{hoy.isoformat()}"
    if con_costos:
        base += "_con_costos"      # así no se confunde con el que no los lleva
    for n in range(1, 50):
        ruta = carpeta / (f"{base}.xlsx" if n == 1 else f"{base}_{n}.xlsx")
        try:
            wb.save(ruta)
            return ruta
        except PermissionError:       # el archivo está abierto en Excel
            continue
    raise ValueError("No se pudo guardar el Excel: cierra los Excel abiertos de "
                     f"la carpeta «{carpeta.name}» e inténtalo otra vez.")